from io import BytesIO

import pytest
from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.security import create_access_token
from app.models.enums import FinishStatus, ModerationStatus, StaffPermission, UserRole
from app.models.group import Group
from app.models.result import Result
from app.services import export_service
from app.services.permissions_service import set_permissions
from tests.factories import make_attendance_with_result, make_event_group, make_user


async def _login(client: AsyncClient, user_id: int) -> None:
    token = create_access_token(user_id)
    resp = await client.get(f"/admin-tools/sso?token={token}", follow_redirects=False)
    assert resp.status_code == 302


def _load(data: bytes) -> Workbook:
    return load_workbook(BytesIO(data))


# --- export_service -----------------------------------------------------


@pytest.mark.asyncio
async def test_build_event_protocol_merges_distance_code_siblings(
    session: AsyncSession,
) -> None:
    org = await make_user(session, "org-rep1@example.com", UserRole.organizer)
    fast = await make_user(session, "fast-rep1@example.com")
    slow = await make_user(session, "slow-rep1@example.com")
    dnf_runner = await make_user(session, "dnf-rep1@example.com")
    event, g1 = await make_event_group(session, org, target_km=33.0)
    g1.distance_code = "X-33"
    g1.name = "X-33 группа #1"
    g2 = Group(
        event_id=event.id,
        location="City",
        name="X-33 группа #2",
        distance_code="X-33",
        target_distance_km=33.0,
    )
    session.add(g2)
    await session.flush()

    await make_attendance_with_result(
        session, g1, fast, finish_status=FinishStatus.finished, moderation=ModerationStatus.approved
    )
    rec_slow = await make_attendance_with_result(
        session, g2, slow, finish_status=FinishStatus.finished, moderation=ModerationStatus.approved
    )
    slow_result = await session.scalar(
        select(Result).where(Result.attendance_record_id == rec_slow.id)
    )
    assert slow_result is not None
    slow_result.duration_seconds = 9000
    await make_attendance_with_result(
        session,
        g2,
        dnf_runner,
        finish_status=FinishStatus.dnf,
        moderation=ModerationStatus.approved,
    )
    await session.commit()

    rows = await export_service.build_event_protocol(session, event)
    assert len(rows) == 3
    assert {r.family_label for r in rows} == {"X-33 · City"}
    finishers = [r for r in rows if r.rank is not None]
    assert [r.rank for r in finishers] == [1, 2]
    # fast (3000s, from the factory default) ranks ahead of slow (9000s)
    # even though slow is in a different pace subgroup of the same family.
    assert finishers[0].duration_seconds == 3000
    assert finishers[1].duration_seconds == 9000
    dnf_rows = [r for r in rows if r.finish_status_label == "DNF"]
    assert len(dnf_rows) == 1
    assert dnf_rows[0].rank is None


@pytest.mark.asyncio
async def test_build_event_protocol_includes_unapproved_as_pending(
    session: AsyncSession,
) -> None:
    """Unlike the public protocol endpoint, the export never hides anything —
    an unapproved self-report still shows up, just unranked, with its
    moderation status visible."""
    org = await make_user(session, "org-rep2@example.com", UserRole.organizer)
    runner = await make_user(session, "runner-rep2@example.com")
    event, group = await make_event_group(session, org)
    rec = await make_attendance_with_result(
        session,
        group,
        runner,
        finish_status=FinishStatus.finished,
        moderation=ModerationStatus.pending,
        self_reported=True,
    )
    await session.commit()

    rows = await export_service.build_event_protocol(session, event)
    assert len(rows) == 1
    assert rows[0].rank is None
    assert rows[0].moderation_label == "на проверке"
    assert rec.id  # sanity: factory actually created it


@pytest.mark.asyncio
async def test_build_results_list_filters_by_date_range(session: AsyncSession) -> None:
    org = await make_user(session, "org-rep3@example.com", UserRole.organizer)
    runner = await make_user(session, "runner-rep3@example.com")
    event_old, group_old = await make_event_group(session, org)
    event_old.date = event_old.date.replace(year=2020)
    event_new, group_new = await make_event_group(session, org)
    await make_attendance_with_result(
        session,
        group_old,
        runner,
        finish_status=FinishStatus.finished,
        moderation=ModerationStatus.approved,
    )
    await make_attendance_with_result(
        session,
        group_new,
        runner,
        finish_status=FinishStatus.finished,
        moderation=ModerationStatus.approved,
    )
    await session.commit()

    all_rows = await export_service.build_results_list(session)
    assert len(all_rows) == 2

    from datetime import date

    recent_rows = await export_service.build_results_list(session, date_from=date(2021, 1, 1))
    assert len(recent_rows) == 1
    assert recent_rows[0].event_title == event_new.title


@pytest.mark.asyncio
async def test_build_summary_scopes_to_organizer(session: AsyncSession) -> None:
    org_a = await make_user(session, "org-rep4a@example.com", UserRole.organizer)
    org_b = await make_user(session, "org-rep4b@example.com", UserRole.organizer)
    runner = await make_user(session, "runner-rep4@example.com")
    event_a, group_a = await make_event_group(session, org_a, target_km=10.0)
    event_b, group_b = await make_event_group(session, org_b, target_km=20.0)
    await make_attendance_with_result(
        session,
        group_a,
        runner,
        finish_status=FinishStatus.finished,
        moderation=ModerationStatus.approved,
    )
    await make_attendance_with_result(
        session,
        group_b,
        runner,
        finish_status=FinishStatus.dnf,
        moderation=ModerationStatus.approved,
    )
    await session.commit()

    everything = await export_service.build_summary(session)
    assert everything.events_count >= 2
    assert everything.finished_count >= 1
    assert everything.dnf_count >= 1

    scoped = await export_service.build_summary(session, created_by=org_a.id)
    assert scoped.events_count == 1
    assert scoped.finished_count == 1
    assert scoped.dnf_count == 0
    assert scoped.total_km == 10.0


# --- HTTP layer -----------------------------------------------------


@pytest.mark.asyncio
async def test_reports_page_requires_permission(session: AsyncSession, client: AsyncClient) -> None:
    org = await make_user(session, "org-rep5@example.com", UserRole.organizer)
    await session.commit()
    await _login(client, org.id)

    resp = await client.get("/admin-tools/reports", follow_redirects=False)
    assert resp.status_code == 302
    assert resp.headers["location"] == "/admin-tools/login"


@pytest.mark.asyncio
async def test_organizer_can_export_own_event_protocol(
    session: AsyncSession, client: AsyncClient
) -> None:
    admin = await make_user(session, "admin-rep6@example.com", UserRole.admin)
    org = await make_user(session, "org-rep6@example.com", UserRole.organizer)
    runner = await make_user(session, "runner-rep6@example.com")
    await set_permissions(session, org, {StaffPermission.reports}, granted_by=admin)
    event, group = await make_event_group(session, org)
    await make_attendance_with_result(
        session,
        group,
        runner,
        finish_status=FinishStatus.finished,
        moderation=ModerationStatus.approved,
    )
    event_id = event.id
    org_id = org.id
    await session.commit()
    await _login(client, org_id)

    resp = await client.get(f"/admin-tools/reports/protocol/{event_id}")
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = _load(resp.content)
    ws = wb["Протокол"]
    assert ws["A3"].value == "Группа"
    assert ws.cell(row=4, column=4).value == f"{runner.first_name} {runner.last_name}"


@pytest.mark.asyncio
async def test_organizer_cannot_export_another_organizers_event(
    session: AsyncSession, client: AsyncClient
) -> None:
    admin = await make_user(session, "admin-rep7@example.com", UserRole.admin)
    owner = await make_user(session, "owner-rep7@example.com", UserRole.organizer)
    other = await make_user(session, "other-rep7@example.com", UserRole.organizer)
    await set_permissions(session, other, {StaffPermission.reports}, granted_by=admin)
    event, _group = await make_event_group(session, owner)
    event_id = event.id
    other_id = other.id
    await session.commit()
    await _login(client, other_id)

    resp = await client.get(f"/admin-tools/reports/protocol/{event_id}", follow_redirects=False)
    assert resp.status_code == 303
    assert resp.headers["location"] == "/admin-tools/reports"


@pytest.mark.asyncio
async def test_summary_report_is_admin_only(session: AsyncSession, client: AsyncClient) -> None:
    admin = await make_user(session, "admin-rep8@example.com", UserRole.admin)
    org = await make_user(session, "org-rep8@example.com", UserRole.organizer)
    await set_permissions(session, org, {StaffPermission.reports}, granted_by=admin)
    org_id = org.id
    admin_id = admin.id
    await session.commit()

    await _login(client, org_id)
    resp = await client.get("/admin-tools/reports/summary", follow_redirects=False)
    assert resp.status_code == 303

    await _login(client, admin_id)
    resp = await client.get("/admin-tools/reports/summary")
    assert resp.status_code == 200
    wb = _load(resp.content)
    ws = wb["Сводка"]
    assert ws["A4"].value == "Показатель"
