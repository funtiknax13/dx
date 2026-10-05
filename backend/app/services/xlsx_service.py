"""Builds the actual .xlsx bytes for /admin-tools/reports from the row data
in app.services.export_service. Kept separate so the query logic there stays
testable without touching openpyxl at all."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.export_service import ProtocolRow, ResultListRow, SummaryStats

_HEADER_FILL = PatternFill("solid", fgColor="0E0E0D")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_CENTER_TOP = Alignment(horizontal="center", vertical="top")


def _fmt_duration(seconds: int | None) -> str:
    if seconds is None:
        return ""
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    return f"{h}:{m:02d}:{s:02d}" if h else f"{m}:{s:02d}"


def _fmt_pace(pace: float | None) -> str:
    if not pace:
        return ""
    m, s = divmod(round(pace), 60)
    return f"{m}:{s:02d}"


def _write_header(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, text in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=text)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _CENTER_TOP
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width


def event_protocol_workbook(
    event_title: str, event_date_label: str, rows: list[ProtocolRow]
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Протокол"
    ws["A1"] = f"{event_title} — {event_date_label}"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:H1")

    headers = [
        "Группа",
        "Подгруппа",
        "Место",
        "Участник",
        "Статус",
        "Модерация",
        "Дистанция, км",
        "Время",
        "Темп, /км",
    ]
    header_row = 3
    _write_header(ws, header_row, headers)

    r = header_row + 1
    merge_start = r
    for i, row in enumerate(rows):
        ws.cell(row=r, column=1, value=row.family_label)
        ws.cell(row=r, column=2, value=row.subgroup_name)
        ws.cell(row=r, column=3, value=row.rank)
        ws.cell(row=r, column=4, value=row.name)
        ws.cell(row=r, column=5, value=row.finish_status_label)
        ws.cell(row=r, column=6, value=row.moderation_label)
        ws.cell(row=r, column=7, value=row.distance_km)
        ws.cell(row=r, column=8, value=_fmt_duration(row.duration_seconds))
        ws.cell(row=r, column=9, value=_fmt_pace(row.pace_seconds_per_km))
        if row.distance_km is not None:
            ws.cell(row=r, column=7).number_format = "0.00"

        next_label = rows[i + 1].family_label if i + 1 < len(rows) else None
        if row.family_label != next_label:
            if r > merge_start:
                ws.merge_cells(start_row=merge_start, start_column=1, end_row=r, end_column=1)
                ws.cell(row=merge_start, column=1).alignment = _CENTER_TOP
            merge_start = r + 1
        r += 1

    _autosize(ws, [22, 20, 7, 26, 12, 12, 13, 10, 10])
    return _save(wb)


def results_list_workbook(title: str, rows: list[ResultListRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:J1")

    headers = [
        "Дата",
        "Событие",
        "Группа",
        "Участник",
        "Статус",
        "Модерация",
        "Дистанция, км",
        "Время",
        "Темп, /км",
        "Источник",
        "Добавлено",
    ]
    header_row = 3
    _write_header(ws, header_row, headers)

    r = header_row + 1
    for row in rows:
        ws.cell(row=r, column=1, value=row.event_date.strftime("%d.%m.%Y"))
        ws.cell(row=r, column=2, value=row.event_title)
        ws.cell(row=r, column=3, value=row.group_name)
        ws.cell(row=r, column=4, value=row.name)
        ws.cell(row=r, column=5, value=row.finish_status_label)
        ws.cell(row=r, column=6, value=row.moderation_label)
        cell = ws.cell(row=r, column=7, value=row.distance_km)
        if row.distance_km is not None:
            cell.number_format = "0.00"
        ws.cell(row=r, column=8, value=_fmt_duration(row.duration_seconds))
        ws.cell(row=r, column=9, value=_fmt_pace(row.pace_seconds_per_km))
        ws.cell(row=r, column=10, value=row.source_label)
        ws.cell(row=r, column=11, value=row.created_at.strftime("%d.%m.%Y %H:%M"))
        r += 1

    _autosize(ws, [11, 24, 16, 26, 12, 12, 13, 10, 10, 12, 16])
    return _save(wb)


def summary_workbook(title: str, stats: SummaryStats) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Сводка"
    ws["A1"] = title
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:B1")
    ws["A2"] = f"Сформировано: {stats.generated_at.strftime('%d.%m.%Y %H:%M')}"

    header_row = 4
    _write_header(ws, header_row, ["Показатель", "Значение"])
    pairs = [
        ("Событий", stats.events_count),
        ("Групп", stats.groups_count),
        ("Всего участий (записей)", stats.attendance_count),
        ("Финишей", stats.finished_count),
        ("DNF", stats.dnf_count),
        ("Уникальных бегунов", stats.unique_runners),
        ("Суммарный километраж (финиши)", round(stats.total_km, 2)),
    ]
    for i, (label, value) in enumerate(pairs, start=header_row + 1):
        ws.cell(row=i, column=1, value=label)
        ws.cell(row=i, column=2, value=value)

    _autosize(ws, [32, 16])
    return _save(wb)


def _save(wb: Workbook) -> bytes:
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
